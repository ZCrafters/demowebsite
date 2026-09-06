"""Generate data/marveile-katalog-audit.xlsx (4 sheets). Honest provenance per row."""
import json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font

OUT = "data/marveile-katalog-audit.xlsx"
TODAY = date(2026, 9, 7).isoformat()
products = json.load(open("data/products.marveile.json", encoding="utf-8"))

wb = Workbook()

# --- Sheet 1: Ringkasan ---
ws = wb.active
ws.title = "Ringkasan"
ws.append(["channel", "url", "jml_terdeteksi", "range_harga", "hero_sku", "status", "tanggal"])
rows = [
    ["Shopee marveile.official", "https://shopee.co.id/marveile.official", 132,
     "179rb-349rb", "Gion/Roux/Gigi/Kylie/Bev", "ok via mirror tokotermurah (langsung butuh JS)", TODAY],
    ["Tokopedia marveile", "https://www.tokopedia.com/marveile", "10+ etalase tampil; toko 18rb terjual",
     "179rb-299rb", "Diora/Oru/Rue/Calla", "ok (fetch parsial, butuh JS utk penuh)", TODAY],
    ["Blibli MAE-70699", "https://www.blibli.com/merchant/marveile/MAE-70699", "kecil (11 ulasan)",
     "-", "-", "ok (toko sepi, bukan acuan utama)", TODAY],
    ["Preloved brand page", "https://preloved.co.id/brands/marveile", 24,
     "40rb-298rb (SECOND, bukan harga resmi)", "Diora/Kai/Lumire", "ok (terpisah dari resmi)", TODAY],
    ["Instagram @marveile.id", "https://www.instagram.com/marveile.id/", 0,
     "-", "Maison Collection", "konten saja (login-block)", TODAY],
    ["TikTok @marveile.id", "https://www.tiktok.com/@marveile.id", 0,
     "-", "live harian 06-21", "konten/live saja (JS-block)", TODAY],
    ["Facebook Marveile.id", "https://www.facebook.com/Marveile.id/", 0,
     "-", "-", "pelengkap (login-block)", TODAY],
    ["Lokal MVP (repo)", "data/products.marveile.json", len(products),
     "provisional", "8 hero + 36 legacy", "terverifikasi lokal", TODAY],
]
for r in rows:
    ws.append(r)

# --- Sheet 2: Produk (resmi saja) ---
wp = wb.create_sheet("Produk")
wp.append(["channel", "nama", "slug", "kategori", "harga", "sizes", "material",
            "warna", "url_gambar", "status_data"])
for p in products:
    wp.append([p.get("source"), p["name"], p["slug"], p.get("category"), p["price"],
               "/".join(p.get("sizes") or []), p.get("material"),
               "/".join(p.get("colors") or []), (p.get("images") or [""])[0],
               "provisional-validated" if not p.get("provisional") else "provisional-observasi"])

# --- Sheet 3: Dedup ---
wd = wb.create_sheet("Dedup")
wd.append(["slug", "muncul_di", "harga_kanonis", "keputusan"])
seen = {}
for p in products:
    seen.setdefault(p["slug"], []).append(f'{p.get("source")}:{p["price"]}')
for slug, hits in sorted(seen.items()):
    prices = sorted(int(h.split(":")[1]) for h in hits)
    wd.append([slug, "; ".join(hits), prices[0],
               "unik" if len(hits) == 1 else "kanonis = termurah"])

# --- Sheet 4: Preloved_Terpisah (second, BUKAN harga resmi) ---
wl = wb.create_sheet("Preloved_Terpisah")
wl.append(["nama_listing", "size", "kondisi", "harga_second", "catatan"])
preloved = [
    ("Kai Knitted Double Zipper", "S", "Sangat baik", 165000, "second"),
    ("Pants Choco", "S", "Sangat baik", 49999, "second"),
    ("Rok Mini", "L", "Sangat baik", 55000, "second"),
    ("Jumpsuit", "S", "Sangat baik", 50000, "second"),
    ("Blouse White", "S", "Sangat baik", 50000, "second"),
    ("Crop Top Black", "S", "Sangat baik", 50000, "second"),
    ("Atasan Biru", "M", "Sangat baik", 80000, "second"),
    ("Jumpsuit S-M", "Other", "Sangat baik", 40000, "second"),
    ("Kemeja Satin Blue", "S", "Sangat baik", 50000, "second"),
    ("Top Blouse", "M", "Sangat baik", 80000, "second"),
    ("Diora Highwaist Pants", "S", "Sangat baik", 65000, "second"),
    ("Elodie Dress", "S", "Sangat baik", 100000, "second"),
    ("Lumire Tweed Blazer", "S", "Baru tanpa tag", 285000, "second"),
    ("Tennis Skort", "M", "Sangat baik", 90000, "second"),
    ("Marvelie One-size", "One size", "Sangat baik", 90000, "second"),
    ("Blazer Black", "S", "Baru tanpa tag", 170000, "second"),
    ("Top Pink", "L", "Baru tanpa tag", 55000, "second"),
    ("Blouse Maroon", "Other", "Sangat baik", 60000, "second"),
    ("Tweed Mini Dress", "S", "Baru dengan tag", 250000, "second"),
    ("Rea Top White", "S", "Sangat baik", 100000, "second"),
    ("Sage Green Blouse", "S", "Sangat baik", 298000, "second"),
    ("Sabrina Top", "S", "Sangat baik", 45000, "second"),
    ("Stacey Dress Polka", "M", "Baru tanpa tag", 285000, "second"),
    ("Joya Cardigan Tanktop Set", "M", "Baru tanpa tag", 215000, "second"),
]
for r in preloved:
    wl.append(list(r) + ["JANGAN dipakai sebagai harga resmi"])

for wsx in (ws, wp, wd, wl):
    wsx.sheet_view.showGridLines = True
    for cell in wsx[1]:
        cell.font = Font(bold=True)

wb.save(OUT)
print(f"wrote {OUT}: Produk={wp.max_row-1} Dedup={wd.max_row-1} Preloved={wl.max_row-1}")
